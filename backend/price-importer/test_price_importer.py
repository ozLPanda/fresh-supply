from __future__ import annotations

import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
import price_importer  # noqa: E402


class PriceImporterTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary_directory.name)

    def tearDown(self) -> None:
        self.temporary_directory.cleanup()

    def save(self, workbook: Workbook, name: str = "price.xlsx") -> Path:
        path = self.root / name
        workbook.save(path)
        workbook.close()
        return path

    def test_extracts_aliases_prices_and_formatted_leading_zero_sku(self) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Котлы"
        worksheet.append(["Прайс-лист"])
        worksheet.append([])
        worksheet.append(
            ["Код", "Наименование", "Розничная цена", "Оптовая цена", "СКО"]
        )
        worksheet.append([1405, "Котёл", "1 234,50", None, 900])
        worksheet["A4"].number_format = "00000000000"
        unrecognized = workbook.create_sheet("Служебный")
        unrecognized.append(["Комментарий", "Значение"])

        result = price_importer.import_price_workbook(self.save(workbook))

        self.assertEqual(result["schemaVersion"], 1)
        self.assertEqual(len(result["rows"]), 1)
        row = result["rows"][0]
        self.assertEqual(row["sourceSheet"], "Котлы")
        self.assertEqual(row["sourceRow"], 4)
        self.assertEqual(row["sku"], "00000001405")
        self.assertEqual(row["name"], "Котёл")
        self.assertEqual(row["price"], "1234.50")
        self.assertIsNone(row["wholesalePrice"])
        self.assertIsNone(row["bulkWholesalePrice"])
        self.assertEqual(row["skoPrice"], "900")
        self.assertEqual(row["errors"], [])
        self.assertEqual(
            result["sheets"][1]["errors"],
            ["unrecognized_sheet:no_supported_header"],
        )

    def test_imports_one_c_mxl_with_sku_from_code_column(self) -> None:
        source = self.root / "one-c-with-code.mxl"
        source.write_bytes(
            b"MOXCEL\x00\x08\x00\x01\x00\x0c\x00\xef\xbb\xbf"
            + """
{16,1,{1,1,{\"#\",\"Код\"}},0},1,
{16,1,{1,1,{\"#\",\"Номенклатура\"}},0},2,
{16,1,{1,1,{\"#\",\"Цена\"}},0},1,0,3,0,
{16,2,{1,1,{\"#\",\"00000000041\"}},0},1,
{16,3,{1,1,{\"#\",\"Строка без цены\"}},0},2,
{16,4,{1,0},0},2,0,3,0,
{16,2,{1,1,{\"#\",\"00000002519\"}},0},1,
{16,3,{1,1,{\"#\",\"Вентиляция вентилятор CDR2E-150\"}},0},2,
{16,4,{1,1,{\"#\",\"29 120,00\"}},0},2,0,3,0,
""".encode("utf-8")
        )

        result = price_importer.import_one_c_mxl(source, "RETAIL")

        self.assertEqual(
            result["rows"],
            [
                {
                    "sourceSheet": "1С MXL",
                    "sourceRow": 2,
                    "sku": "00000000041",
                    "name": "Строка без цены",
                    "price": None,
                    "wholesalePrice": None,
                    "bulkWholesalePrice": None,
                    "skoPrice": None,
                    "errors": [],
                    "missingRetailPrice": True,
                },
                {
                    "sourceSheet": "1С MXL",
                    "sourceRow": 3,
                    "sku": "00000002519",
                    "name": "Вентиляция вентилятор CDR2E-150",
                    "price": "29120.00",
                    "wholesalePrice": None,
                    "bulkWholesalePrice": None,
                    "skoPrice": None,
                    "errors": [],
                }
            ],
        )
        self.assertEqual(result["sheets"][0]["columns"], {"sku": 1, "name": 2, "price": 3})

    def test_imports_one_c_mxl_with_doubled_quotes_in_product_name(self) -> None:
        source = self.root / "one-c-with-quoted-name.mxl"
        source.write_bytes(
            b"MOXCEL\x00\x08\x00\x01\x00\x0c\x00\xef\xbb\xbf"
            + """
{16,1,{1,1,{"#","Код"}},0},1,
{16,1,{1,1,{"#","Номенклатура"}},0},2,
{16,1,{1,1,{"#","Цена"}},0},1,0,3,0,
{16,2,{1,1,{"#","00000001614"}},0},1,
{16,3,{1,1,{"#","Кран шаровый (STA) 1""(25) PN25 син.ручка вн/нар (арт-1614)"}},0},2,
{16,4,{1,1,{"#","5 328,00"}},0},2,0,3,0,
""".encode("utf-8")
        )

        result = price_importer.import_one_c_mxl(source, "RETAIL")

        self.assertEqual(
            result["rows"],
            [
                {
                    "sourceSheet": "1С MXL",
                    "sourceRow": 2,
                    "sku": "00000001614",
                    "name": 'Кран шаровый (STA) 1"(25) PN25 син.ручка вн/нар (арт-1614)',
                    "price": "5328.00",
                    "wholesalePrice": None,
                    "bulkWholesalePrice": None,
                    "skoPrice": None,
                    "errors": [],
                }
            ],
        )

    def test_treats_zero_as_missing_and_rejects_negative_or_non_finite_prices(self) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["Артикул", "Название", "Цена", "Опт", "Крупный опт"])
        worksheet.append(["A-1", "Первый", -1, 0, "NaN"])
        worksheet.append(["A-2", "Второй", "Infinity", 10, 20])

        result = price_importer.import_price_workbook(self.save(workbook))

        first, second = result["rows"]
        self.assertIsNone(first["price"])
        self.assertIsNone(first["wholesalePrice"])
        self.assertIsNone(first["bulkWholesalePrice"])
        self.assertCountEqual(
            first["errors"],
            ["invalid_price", "invalid_bulk_wholesale_price"],
        )
        self.assertIsNone(second["price"])
        self.assertIn("invalid_price", second["errors"])

    def test_marks_every_normalized_duplicate_sku(self) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["SKU", "Товар", "Розница"])
        worksheet.append([" A-001 ", "Первый", 100])
        worksheet.append(["a-001", "Второй", 200])

        result = price_importer.import_price_workbook(self.save(workbook))

        self.assertEqual([row["sku"] for row in result["rows"]], ["A-001", "a-001"])
        self.assertTrue(all("duplicate_sku" in row["errors"] for row in result["rows"]))

    def test_collapses_duplicate_sku_when_all_prices_are_compatible(self) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["SKU", "Товар", "Розница", "Опт"])
        worksheet.append([" A-001 ", "Первый", 100, None])
        worksheet.append(["a-001", "Повтор", 100, 90])

        result = price_importer.import_price_workbook(self.save(workbook))

        self.assertEqual(len(result["rows"]), 1)
        self.assertEqual(result["rows"][0]["sku"], "A-001")
        self.assertEqual(result["rows"][0]["price"], "100")
        self.assertEqual(result["rows"][0]["wholesalePrice"], "90")
        self.assertEqual(result["rows"][0]["errors"], [])

    def test_recognizes_dymohody_heading_as_product_name_column(self) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["Код", None, "Дымоходы Т и С", "Розница"])
        worksheet.append(["1718", None, "Конденсатоотвод Термо", 3908])

        result = price_importer.import_price_workbook(self.save(workbook))

        self.assertEqual(len(result["rows"]), 1)
        self.assertEqual(result["rows"][0]["name"], "Конденсатоотвод Термо")
        self.assertEqual(result["rows"][0]["errors"], [])

    def test_skips_rows_where_all_prices_are_zero(self) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["SKU", "Цена", "Опт"])
        worksheet.append(["A-1", 0, 0])

        result = price_importer.import_price_workbook(self.save(workbook))

        self.assertEqual(result["rows"], [])

    def test_recognizes_price_only_sheet_and_keeps_name_optional(self) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["Артикул", "Опт"])
        worksheet.append(["A-1", 125])

        result = price_importer.import_price_workbook(self.save(workbook))

        self.assertTrue(result["sheets"][0]["recognized"])
        self.assertNotIn("name", result["sheets"][0]["columns"])
        self.assertIsNone(result["rows"][0]["name"])
        self.assertIsNone(result["rows"][0]["price"])
        self.assertEqual(result["rows"][0]["wholesalePrice"], "125")
        self.assertEqual(result["rows"][0]["errors"], [])

    def test_recognizes_raw_catalog_headers_and_ignores_formatted_far_columns(self) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Исходный каталог"
        worksheet.append([None, "ТОО GastroFlow"])
        worksheet.append(
            [
                "Код",
                "Изображение",
                "Наименование",
                "Ед. изм",
                "Остаток",
                "Розница (тг)",
                "*Для посотянных клиентов",
                "Более 350 000",
            ]
        )
        worksheet.append(
            ["00000001405", None, "Расширительный бак", "шт", "в наличии", 35200, 28161, 22489]
        )
        worksheet.cell(row=1, column=99, value="служебное форматирование")

        with patch.object(price_importer, "MIN_WORKSHEET_XML_OPTIMIZE_BYTES", 0):
            result = price_importer.import_price_workbook(self.save(workbook))

        self.assertTrue(result["sheets"][0]["recognized"])
        self.assertIn("column_limit_exceeded:99>50", result["sheets"][0]["errors"])
        self.assertEqual(len(result["rows"]), 1)
        row = result["rows"][0]
        self.assertEqual(row["sku"], "00000001405")
        self.assertEqual(row["price"], "35200")
        self.assertEqual(row["wholesalePrice"], "28161")
        self.assertEqual(row["bulkWholesalePrice"], "22489")
        self.assertEqual(row["errors"], [])

    def test_propagates_duplicate_header_error_to_every_data_row(self) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["SKU", "Код", "Цена"])
        worksheet.append(["A-1", "IGNORED-1", 100])
        worksheet.append(["A-2", "IGNORED-2", 200])

        result = price_importer.import_price_workbook(self.save(workbook))

        self.assertIn("duplicate_header:sku", result["sheets"][0]["errors"])
        self.assertTrue(
            all("duplicate_header:sku" in row["errors"] for row in result["rows"])
        )

    def test_searches_headers_only_in_first_thirty_rows(self) -> None:
        workbook = Workbook()
        recognized = workbook.active
        recognized.title = "Row 30"
        for _ in range(29):
            recognized.append([""])
        recognized.append(["Код", "Наименование", "Цена"])
        recognized.append(["A-1", "Товар", 100])

        too_late = workbook.create_sheet("Row 31")
        for _ in range(30):
            too_late.append([""])
        too_late.append(["Код", "Наименование", "Цена"])
        too_late.append(["A-2", "Товар", 200])

        result = price_importer.import_price_workbook(self.save(workbook))

        self.assertTrue(result["sheets"][0]["recognized"])
        self.assertEqual(result["sheets"][0]["headerRow"], 30)
        self.assertFalse(result["sheets"][1]["recognized"])
        self.assertEqual(len(result["rows"]), 1)

    def test_enforces_row_sheet_column_and_cell_limits(self) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["Код", "Название", "Цена"])
        worksheet.append(["A-1", "X" * 501, 100])
        worksheet.append(["A-2", "Второй", 200])
        worksheet.append(["A-3", "Третий", 300])

        with patch.object(price_importer, "MAX_DATA_ROWS", 2):
            result = price_importer.import_price_workbook(self.save(workbook))

        self.assertEqual(len(result["rows"]), 2)
        self.assertIn("cell_too_long:name", result["rows"][0]["errors"])
        self.assertEqual(len(result["rows"][0]["name"]), 500)
        self.assertIn("row_limit_exceeded:3>2", result["errors"])

        many_sheets = Workbook()
        many_sheets.create_sheet("Second")
        with patch.object(price_importer, "MAX_SHEETS", 1):
            limited = price_importer.import_price_workbook(
                self.save(many_sheets, "many-sheets.xlsx")
            )
        self.assertEqual(limited["rows"], [])
        self.assertEqual(len(limited["sheets"]), 2)
        self.assertIn("sheet_limit_exceeded:2>1", limited["errors"])

        many_columns = Workbook()
        wide = many_columns.active
        wide.append(["Код", "Название", "Цена"])
        wide.append(["A-1", "Товар", 100])
        wide.cell(row=1, column=51, value="Служебная область")
        with patch.object(price_importer, "MAX_COLUMNS", 50):
            wide_result = price_importer.import_price_workbook(
                self.save(many_columns, "wide.xlsx")
            )
        self.assertIn("column_limit_exceeded:51>50", wide_result["sheets"][0]["errors"])
        self.assertEqual(wide_result["rows"][0]["sku"], "A-1")

    def test_writes_json_atomically_and_cli_handles_success_and_bad_extension(self) -> None:
        output = self.root / "nested" / "result.json"
        payload = {"schemaVersion": 1, "sheets": [], "rows": [], "errors": []}

        price_importer.write_result(output, payload)

        self.assertEqual(json.loads(output.read_text(encoding="utf-8")), payload)

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["SKU", "Цена"])
        worksheet.append(["A-1", 100])
        source = self.save(workbook, "cli.xlsx")
        cli_output = self.root / "cli.json"
        self.assertEqual(price_importer.main([str(source), str(cli_output)]), 0)
        self.assertEqual(
            json.loads(cli_output.read_text(encoding="utf-8"))["rows"][0]["sku"],
            "A-1",
        )

        unsupported = self.root / "price.xls"
        unsupported.write_bytes(b"not an xlsx")
        self.assertEqual(price_importer.main([str(unsupported), str(output)]), 2)


if __name__ == "__main__":
    unittest.main()
